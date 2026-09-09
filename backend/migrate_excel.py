import os
import openpyxl
from datetime import datetime
from backend.database import SessionLocal, engine, Base
from backend.models import Articulo, Cliente, Factura, DetalleFactura, Compra, Despacho, DetalleDespacho, Configuracion, Usuario
from backend.auth import hash_password

EXCEL_PATH = r"C:\Users\e\Downloads\Prolago Carora 2026.xlsm"

def run_migration():
    print("Creando tablas en la base de datos...")
    Base.metadata.create_all(bind=engine)
    db = SessionLocal()

    try:
        # 1. Configuración inicial
        print("Configurando parámetros iniciales...")
        config_items = {
            "nombre_empresa": "Prolago Carora",
            "rif_empresa": "J-41215654-3",
            "direccion_empresa": "Av. Rotaria, Carora, Edo. Lara",
            "telefono_empresa": "(0252) 421-XXXX",
            "tasa_bcv": "473.92"
        }
        for clave, valor in config_items.items():
            existing = db.query(Configuracion).filter_by(clave=clave).first()
            if not existing:
                db.add(Configuracion(clave=clave, valor=valor))
        
        # 2. Usuarios por defecto
        admin_user = db.query(Usuario).filter_by(username="admin").first()
        if not admin_user:
            db.add(Usuario(
                username="admin",
                password_hash=hash_password("admin123"),
                nombre="Administrador General",
                rol="admin",
                activo=True
            ))
            db.add(Usuario(
                username="cajero",
                password_hash=hash_password("cajero123"),
                nombre="Caja Principal",
                rol="cajero",
                activo=True
            ))
            print("Usuarios creados: admin / admin123  y  cajero / cajero123")

        db.commit()

        if not os.path.exists(EXCEL_PATH):
            print(f"No se encontró el archivo Excel en {EXCEL_PATH}")
            return

        print(f"Cargando {EXCEL_PATH} (esto puede tardar unos segundos)...")
        wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
        
        # Leer Tasa BCV real de la hoja Compras si existe
        if "Compras" in wb.sheetnames:
            ws_comp = wb["Compras"]
            tasa_val = ws_comp["K3"].value
            if tasa_val and isinstance(tasa_val, (int, float)):
                tasa_obj = db.query(Configuracion).filter_by(clave="tasa_bcv").first()
                if tasa_obj:
                    tasa_obj.valor = str(round(float(tasa_val), 2))
                    db.commit()
                    print(f"Tasa BCV importada desde Excel: {tasa_obj.valor} Bs./$")

        # 3. Migrar Clientes
        clientes_agregados = 0
        if "Clientes" in wb.sheetnames:
            ws_cli = wb["Clientes"]
            for r in range(5, ws_cli.max_row + 1):
                ci = str(ws_cli.cell(r, 2).value or "").strip()
                nombre = str(ws_cli.cell(r, 3).value or "").strip()
                direc = str(ws_cli.cell(r, 4).value or "").strip()
                telf = str(ws_cli.cell(r, 5).value or "").strip()
                if nombre and nombre != "0" and len(nombre) > 2:
                    exists = db.query(Cliente).filter_by(nombre=nombre).first()
                    if not exists:
                        db.add(Cliente(
                            cedula_rif=ci,
                            nombre=nombre,
                            direccion=direc,
                            telefono=telf,
                            tipo="general"
                        ))
                        clientes_agregados += 1

        if "Clientes2" in wb.sheetnames:
            ws_cli2 = wb["Clientes2"]
            for r in range(6, ws_cli2.max_row + 1):
                nombre = str(ws_cli2.cell(r, 3).value or "").strip()
                direc = str(ws_cli2.cell(r, 4).value or "").strip()
                if nombre and len(nombre) > 2:
                    exists = db.query(Cliente).filter_by(nombre=nombre).first()
                    if not exists:
                        db.add(Cliente(
                            cedula_rif="",
                            nombre=nombre,
                            direccion=direc,
                            telefono="",
                            tipo="despacho"
                        ))
                        clientes_agregados += 1
        
        db.commit()
        print(f"Clientes migrados: {clientes_agregados}")

        # 4. Migrar Artículos (Inventario y Catálogo)
        articulos_sheet = None
        for s in wb.sheetnames:
            if "Art" in s:
                articulos_sheet = wb[s]
                break

        articulos_agregados = 0
        if articulos_sheet:
            # Columnas esperadas según inspección:
            # Col A(1): Código
            # Col B(2): Artículo (Nombre)
            # Col C(3): Entrada
            # Col D(4): Salida
            # Col E(5): Stock
            # Col F(6): Costo
            # Col G(7): Flete
            # Col H(8): Costo Final
            # Col I(9): mas
            # Col J(10): Rent.
            # Col K(11): Sugerido
            # Col L(12): Precio
            # Col M(13): Categoría
            # Col N(14): Marca
            # Col O(15): Proveedor
            for r in range(5, articulos_sheet.max_row + 1):
                raw_codigo = articulos_sheet.cell(r, 1).value
                if raw_codigo is None:
                    continue
                try:
                    codigo = int(float(str(raw_codigo).strip()))
                except ValueError:
                    continue
                
                nombre = str(articulos_sheet.cell(r, 2).value or "").strip()
                if not nombre:
                    continue
                
                def to_flt(val):
                    if val is None or val == "" or str(val).startswith("#"):
                        return 0.0
                    try:
                        return float(val)
                    except:
                        return 0.0

                stock = to_flt(articulos_sheet.cell(r, 5).value)
                costo = to_flt(articulos_sheet.cell(r, 6).value)
                flete = to_flt(articulos_sheet.cell(r, 7).value)
                costo_final = to_flt(articulos_sheet.cell(r, 8).value) or (costo + flete)
                mas = to_flt(articulos_sheet.cell(r, 9).value)
                rent = to_flt(articulos_sheet.cell(r, 10).value)
                sugerido = to_flt(articulos_sheet.cell(r, 11).value)
                precio = to_flt(articulos_sheet.cell(r, 12).value)
                cat = str(articulos_sheet.cell(r, 13).value or "GENERAL").strip()
                marca = str(articulos_sheet.cell(r, 14).value or "").strip()
                prov = str(articulos_sheet.cell(r, 15).value or "").strip()

                exists = db.query(Articulo).filter_by(codigo=codigo).first()
                if not exists:
                    db.add(Articulo(
                        codigo=codigo,
                        nombre=nombre,
                        categoria=cat if cat != "None" else "GENERAL",
                        marca=marca if marca != "None" else "",
                        proveedor=prov if prov != "None" else "",
                        costo=costo,
                        flete=flete,
                        costo_final=costo_final,
                        mas=mas,
                        rentabilidad=rent,
                        sugerido=sugerido,
                        precio=precio,
                        stock=stock,
                        activo=True
                    ))
                    articulos_agregados += 1
            
            db.commit()
            print(f"Artículos migrados exitosamente: {articulos_agregados}")

        print("\n¡Migración inicial completada con éxito!")

    except Exception as e:
        db.rollback()
        print("Error durante la migración:", e)
        raise
    finally:
        db.close()

if __name__ == "__main__":
    run_migration()
