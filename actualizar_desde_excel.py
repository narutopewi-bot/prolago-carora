import os
import shutil
import openpyxl
from datetime import datetime
from backend.database import SessionLocal, engine, Base
from backend.models import Articulo, Configuracion

EXCEL_PATH = r"C:\Users\e\Documents\Prolago Carora 2026.xlsm"
DB_PATH = os.path.abspath("./prolago.db")

def actualizar_base_datos():
    print(f"=== INICIANDO ACTUALIZACIÓN DESDE EXCEL ===")
    print(f"Archivo Excel: {EXCEL_PATH}")
    print(f"Base de datos destino: {DB_PATH}")

    if not os.path.exists(EXCEL_PATH):
        print(f"ERROR: No se encontró el archivo {EXCEL_PATH}")
        return

    # 1. Crear copia de respaldo preventiva de prolago.db
    backup_dir = os.path.abspath("./respaldos")
    os.makedirs(backup_dir, exist_ok=True)
    fecha_str = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup_path = os.path.join(backup_dir, f"prolago_backup_antes_excel_{fecha_str}.db")
    if os.path.exists(DB_PATH):
        shutil.copy2(DB_PATH, backup_path)
        print(f"Copia de seguridad preventiva creada: {backup_path}")

    # 2. Cargar Excel
    print("Cargando libro de Excel con openpyxl (esto toma unos segundos)...")
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)

    ws_art = None
    for s in wb.sheetnames:
        if "art" in s.lower():
            ws_art = wb[s]
            print(f"Hoja de artículos identificada: '{s}' con {ws_art.max_row} filas")
            break

    if not ws_art:
        print("ERROR: No se encontró hoja de artículos en el Excel.")
        return

    db = SessionLocal()
    try:
        arts_db = {a.codigo: a for a in db.query(Articulo).all()}
        print(f"Artículos existentes en base de datos: {len(arts_db)}")

        def to_flt(val):
            if val is None or val == "" or str(val).startswith("#"):
                return 0.0
            try:
                return float(val)
            except:
                return 0.0

        nuevos_creados = 0
        precios_actualizados = 0
        stocks_actualizados = 0
        costos_actualizados = 0
        total_procesados = 0

        lista_precios_cambiados = []
        lista_articulos_nuevos = []

        # Recorrer artículos desde la fila 5
        for r in range(5, ws_art.max_row + 1):
            raw_cod = ws_art.cell(r, 1).value
            if raw_cod is None:
                continue
            try:
                codigo = int(float(str(raw_cod).strip()))
            except ValueError:
                continue

            nombre = str(ws_art.cell(r, 2).value or "").strip()
            if not nombre:
                continue

            total_procesados += 1

            stock = to_flt(ws_art.cell(r, 5).value)
            costo = to_flt(ws_art.cell(r, 6).value)
            flete = to_flt(ws_art.cell(r, 7).value)
            costo_final = to_flt(ws_art.cell(r, 8).value) or round(costo + flete, 2)
            mas = to_flt(ws_art.cell(r, 9).value)
            rent = to_flt(ws_art.cell(r, 10).value)
            sugerido = to_flt(ws_art.cell(r, 11).value)
            precio = to_flt(ws_art.cell(r, 12).value)
            
            cat = str(ws_art.cell(r, 13).value or "GENERAL").strip()
            if cat == "None" or not cat:
                cat = "GENERAL"
            marca = str(ws_art.cell(r, 14).value or "").strip()
            if marca == "None":
                marca = ""
            prov = str(ws_art.cell(r, 15).value or "").strip()
            if prov == "None":
                prov = ""

            if codigo not in arts_db:
                # ARTÍCULO NUEVO
                nuevo_art = Articulo(
                    codigo=codigo,
                    nombre=nombre,
                    categoria=cat,
                    marca=marca,
                    proveedor=prov,
                    costo=costo,
                    flete=flete,
                    costo_final=costo_final,
                    mas=mas,
                    rentabilidad=rent,
                    sugerido=sugerido,
                    precio=precio,
                    stock=stock,
                    stock_alerta=5.0,
                    activo=True
                )
                db.add(nuevo_art)
                arts_db[codigo] = nuevo_art
                nuevos_creados += 1
                lista_articulos_nuevos.append((codigo, nombre, precio, stock, cat))
            else:
                # ARTÍCULO EXISTENTE -> ACTUALIZAR
                art = arts_db[codigo]
                
                # Chequear si cambió el precio
                precio_anterior = art.precio or 0.0
                if abs(precio_anterior - precio) > 0.005:
                    lista_precios_cambiados.append((codigo, nombre, precio_anterior, precio))
                    precios_actualizados += 1

                # Chequear si cambió el stock
                stock_anterior = art.stock or 0.0
                if abs(stock_anterior - stock) > 0.005:
                    stocks_actualizados += 1

                # Chequear si cambiaron costos
                if abs((art.costo or 0.0) - costo) > 0.005 or abs((art.costo_final or 0.0) - costo_final) > 0.005:
                    costos_actualizados += 1

                # Aplicar actualizaciones
                art.nombre = nombre
                art.categoria = cat
                art.marca = marca
                art.proveedor = prov
                art.costo = costo
                art.flete = flete
                art.costo_final = costo_final
                art.mas = mas
                art.rentabilidad = rent
                art.sugerido = sugerido
                art.precio = precio
                art.stock = stock
                art.activo = True

        db.commit()
        print("\n--- RESUMEN DE ACTUALIZACIÓN ---")
        print(f"Total artículos procesados del Excel: {total_procesados}")
        print(f"Artículos NUEVOS agregados: {nuevos_creados}")
        print(f"Artículos con PRECIO modificado y actualizado: {precios_actualizados}")
        print(f"Artículos con STOCK actualizado: {stocks_actualizados}")
        print(f"Artículos con COSTOS/RENTABILIDAD actualizados: {costos_actualizados}")
        print(f"Total artículos en la base de datos ahora: {db.query(Articulo).count()}")

        # 3. Sincronizar copias en carpetas dist si existen
        for dist_dest in ["./dist/prolago.db", "./dist/ProlagoCarora/_internal/prolago.db"]:
            if os.path.exists(os.path.dirname(dist_dest)):
                shutil.copy2(DB_PATH, dist_dest)
                print(f"Copia sincronizada en: {dist_dest}")

        return {
            "total_procesados": total_procesados,
            "nuevos_creados": nuevos_creados,
            "precios_actualizados": precios_actualizados,
            "stocks_actualizados": stocks_actualizados,
            "costos_actualizados": costos_actualizados,
            "lista_precios_cambiados": lista_precios_cambiados,
            "lista_articulos_nuevos": lista_articulos_nuevos
        }

    except Exception as e:
        db.rollback()
        print(f"ERROR durante la actualización: {e}")
        raise
    finally:
        db.close()

if __name__ == "__main__":
    actualizar_base_datos()
